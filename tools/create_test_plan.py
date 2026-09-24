from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / "docs"
OUTPUT_FILE = DOCS_DIR / "TEST_PLAN_60_CASES.xlsx"

DOCS_DIR.mkdir(parents=True, exist_ok=True)


cases = [
    # M01 - Authentication & Roles: 7
    ("TC001", "M01 Authentication & Roles", "NORMAL",
     "Đăng nhập bằng tài khoản hợp lệ",
     "Tài khoản tồn tại và đang hoạt động",
     "Nhập email và mật khẩu đúng; gửi yêu cầu đăng nhập",
     "Đăng nhập thành công và tạo phiên hợp lệ",
     "High"),

    ("TC002", "M01 Authentication & Roles", "NORMAL",
     "Đăng xuất khỏi hệ thống",
     "Đang đăng nhập",
     "Chọn chức năng đăng xuất",
     "Phiên bị hủy và trang yêu cầu đăng nhập không truy cập lại được",
     "High"),

    ("TC003", "M01 Authentication & Roles", "VALIDATION",
     "Đăng nhập thiếu email",
     "Không yêu cầu tài khoản đặc biệt",
     "Để trống email; nhập mật khẩu; gửi",
     "Backend trả 422 và thông báo lỗi hợp lệ",
     "Medium"),

    ("TC004", "M01 Authentication & Roles", "VALIDATION",
     "Đăng ký email sai định dạng",
     "Trang đăng ký hoạt động",
     "Nhập email không hợp lệ",
     "Backend trả 422",
     "Medium"),

    ("TC005", "M01 Authentication & Roles", "AUTHORIZATION",
     "Guest truy cập endpoint dành cho user",
     "Chưa đăng nhập",
     "Gọi GET /api/v1/me/collections",
     "Backend trả 401",
     "High"),

    ("TC006", "M01 Authentication & Roles", "AUTHORIZATION",
     "User truy cập endpoint admin",
     "Đăng nhập user thường",
     "Gọi trực tiếp endpoint quản trị",
     "Backend trả 403 và không lộ dữ liệu quản trị",
     "Critical"),

    ("TC007", "M01 Authentication & Roles", "AUTHORIZATION",
     "Sử dụng lại phiên sau khi đăng xuất",
     "Đã đăng xuất",
     "Gửi lại request bằng phiên cũ",
     "Request bị từ chối",
     "High"),

    # M02 - Museum Spaces & Panoramas: 7
    ("TC008", "M02 Museum Spaces & Panoramas", "NORMAL",
     "Mở danh sách không gian",
     "Có dữ liệu museum_spaces",
     "Mở /tham-quan-360",
     "Danh sách không gian được hiển thị",
     "Medium"),

    ("TC009", "M02 Museum Spaces & Panoramas", "NORMAL",
     "Mở chi tiết không gian",
     "Space id hợp lệ",
     "Gọi GET /api/v1/spaces/1",
     "Trả về dữ liệu space tương ứng",
     "Medium"),

    ("TC010", "M02 Museum Spaces & Panoramas", "NORMAL",
     "Lấy panorama theo không gian",
     "Space tồn tại",
     "Gọi endpoint panorama của space",
     "Trả về danh sách panorama đúng quan hệ",
     "High"),

    ("TC011", "M02 Museum Spaces & Panoramas", "VALIDATION",
     "Gọi space với id không hợp lệ kiểu chữ",
     "API đang hoạt động",
     "Gọi /api/v1/spaces/abc",
     "Request bị từ chối hoặc route không hợp lệ; không 500",
     "Medium"),

    ("TC012", "M02 Museum Spaces & Panoramas", "VALIDATION",
     "Gọi panorama với id bằng 0",
     "API đang hoạt động",
     "Gọi /api/v1/panoramas/0",
     "Không lấy dữ liệu không hợp lệ",
     "Medium"),

    ("TC013", "M02 Museum Spaces & Panoramas", "AUTHORIZATION",
     "Truy cập dữ liệu space quản trị bằng endpoint sửa đổi",
     "User thường đăng nhập",
     "Gửi request sửa space",
     "Backend trả 403",
     "High"),

    ("TC014", "M02 Museum Spaces & Panoramas", "AUTHORIZATION",
     "Truy cập space không tồn tại",
     "API hoạt động",
     "Gọi space id rất lớn",
     "Backend trả 404 và không lộ SQL",
     "Medium"),

    # M03 - Artifacts & Hotspots: 7
    ("TC015", "M03 Artifacts & Hotspots", "NORMAL",
     "Mở thông tin hiện vật",
     "Artifact tồn tại",
     "Gọi GET /api/v1/artifacts/1",
     "Hiển thị thông tin hiện vật",
     "High"),

    ("TC016", "M03 Artifacts & Hotspots", "NORMAL",
     "Lấy hotspot theo panorama",
     "Panorama có hotspot",
     "Gọi endpoint hotspots",
     "Các hotspot được trả về",
     "High"),

    ("TC017", "M03 Artifacts & Hotspots", "NORMAL",
     "Mở chi tiết hotspot",
     "Hotspot tồn tại",
     "Gọi GET /api/v1/hotspots/1",
     "Trả về thông tin hotspot",
     "High"),

    ("TC018", "M03 Artifacts & Hotspots", "VALIDATION",
     "Hotspot có tọa độ không hợp lệ",
     "Có quyền quản trị dữ liệu test",
     "Gửi yaw/pitch ngoài phạm vi schema",
     "Backend từ chối dữ liệu",
     "High"),

    ("TC019", "M03 Artifacts & Hotspots", "VALIDATION",
     "Hotspot tham chiếu artifact không tồn tại",
     "Có quyền quản trị dữ liệu test",
     "Gửi artifact_id không tồn tại",
     "Foreign key hoặc validation từ chối",
     "High"),

    ("TC020", "M03 Artifacts & Hotspots", "AUTHORIZATION",
     "User thường sửa hotspot",
     "Đăng nhập user thường",
     "Gọi endpoint quản trị hotspot",
     "Trả 403",
     "High"),

    ("TC021", "M03 Artifacts & Hotspots", "AUTHORIZATION",
     "Gọi artifact không tồn tại",
     "API hoạt động",
     "Gọi artifact id lớn",
     "Trả 404",
     "Medium"),

    # M04 - Transitions: 6
    ("TC022", "M04 360 Scene Transitions", "NORMAL",
     "Lấy transition của panorama",
     "Panorama có transition",
     "Gọi endpoint transitions",
     "Danh sách transition được trả về",
     "High"),

    ("TC023", "M04 360 Scene Transitions", "NORMAL",
     "Thực hiện chuyển panorama",
     "Source và target tồn tại",
     "Chọn transition trên viewer",
     "Viewer chuyển sang panorama đích",
     "Critical"),

    ("TC024", "M04 360 Scene Transitions", "NORMAL",
     "Chuyển nhiều scene liên tiếp",
     "Có chuỗi transition hợp lệ",
     "Thực hiện 3 chuyển cảnh",
     "Không mất trạng thái panorama",
     "High"),

    ("TC025", "M04 360 Scene Transitions", "VALIDATION",
     "Transition trỏ tới panorama không tồn tại",
     "Có quyền quản trị dữ liệu",
     "Tạo transition với target id sai",
     "Backend từ chối",
     "High"),

    ("TC026", "M04 360 Scene Transitions", "VALIDATION",
     "Transition có tọa độ không hợp lệ",
     "Có quyền quản trị dữ liệu",
     "Gửi pitch/yaw sai miền",
     "Validation từ chối",
     "Medium"),

    ("TC027", "M04 360 Scene Transitions", "AUTHORIZATION",
     "User thường tạo transition",
     "Đăng nhập user",
     "POST vào endpoint quản trị transition",
     "Trả 403",
     "High"),

    # M05 - Audio: 6
    ("TC028", "M05 Multilingual Audio", "NORMAL",
     "Lấy danh sách ngôn ngữ",
     "Có dữ liệu languages",
     "Gọi GET /api/v1/languages",
     "Danh sách ngôn ngữ trả về",
     "Medium"),

    ("TC029", "M05 Multilingual Audio", "NORMAL",
     "Lấy audio narration theo artifact",
     "Artifact có audio",
     "Gọi endpoint audio",
     "Trả về narration tương ứng",
     "High"),

    ("TC030", "M05 Multilingual Audio", "NORMAL",
     "Lọc narration theo ngôn ngữ",
     "Có nhiều language",
     "Gửi language_id hợp lệ",
     "Trả về narration đúng ngôn ngữ",
     "High"),

    ("TC031", "M05 Multilingual Audio", "VALIDATION",
     "language_id không tồn tại",
     "API hoạt động",
     "Gửi language_id sai",
     "Backend trả 404 hoặc 422 theo contract",
     "Medium"),

    ("TC032", "M05 Multilingual Audio", "VALIDATION",
     "Audio narration thiếu quan hệ bắt buộc",
     "Có quyền quản trị dữ liệu",
     "Gửi dữ liệu thiếu trường bắt buộc",
     "Backend trả 422",
     "High"),

    ("TC033", "M05 Multilingual Audio", "AUTHORIZATION",
     "User thường sửa narration",
     "Đăng nhập user",
     "Gọi endpoint quản trị narration",
     "Trả 403",
     "High"),

    # M06 - Tours & Modes: 7
    ("TC034", "M06 Themed Tours & Modes", "NORMAL",
     "Lấy danh sách tour",
     "Có dữ liệu tours",
     "Gọi GET /api/v1/tours",
     "Danh sách tour hiển thị",
     "High"),

    ("TC035", "M06 Themed Tours & Modes", "NORMAL",
     "Lọc tour free",
     "Có tour free",
     "Gửi mode=free",
     "Chỉ trả tour mode free",
     "Medium"),

    ("TC036", "M06 Themed Tours & Modes", "NORMAL",
     "Lọc tour guided",
     "Có tour guided",
     "Gửi mode=guided",
     "Chỉ trả tour mode guided",
     "Medium"),

    ("TC037", "M06 Themed Tours & Modes", "NORMAL",
     "Lấy stop của tour",
     "Tour tồn tại",
     "Gọi /tours/{id}/stops",
     "Stop có stop_order đúng",
     "High"),

    ("TC038", "M06 Themed Tours & Modes", "VALIDATION",
     "mode không thuộc free/guided",
     "API hoạt động",
     "Gửi mode=abc",
     "Backend trả 422",
     "Medium"),

    ("TC039", "M06 Themed Tours & Modes", "AUTHORIZATION",
     "User sửa tour",
     "Đăng nhập user thường",
     "Gọi endpoint sửa tour",
     "Trả 403",
     "High"),

    ("TC040", "M06 Themed Tours & Modes", "AUTHORIZATION",
     "User sửa tour stop",
     "Đăng nhập user thường",
     "Gọi endpoint quản trị stop",
     "Trả 403",
     "High"),

    # M07 - Guestbook & Collection: 7
    ("TC041", "M07 Guestbook & Collection", "NORMAL",
     "Xem guestbook",
     "Có entry được phép hiển thị",
     "Mở guestbook",
     "Các entry hợp lệ hiển thị",
     "Medium"),

    ("TC042", "M07 Guestbook & Collection", "NORMAL",
     "User gửi guestbook",
     "User đã đăng nhập",
     "Nhập nội dung hợp lệ và gửi",
     "Entry được tạo",
     "High"),

    ("TC043", "M07 Guestbook & Collection", "NORMAL",
     "Thêm artifact vào collection",
     "User đăng nhập và artifact tồn tại",
     "POST artifact_id",
     "Collection được tạo",
     "High"),

    ("TC044", "M07 Guestbook & Collection", "NORMAL",
     "Xem collection cá nhân",
     "User có collection",
     "Mở collection",
     "Chỉ collection của tài khoản hiện tại được trả về",
     "Critical"),

    ("TC045", "M07 Guestbook & Collection", "VALIDATION",
     "Thêm collection với artifact_id không hợp lệ",
     "User đăng nhập",
     "Gửi artifact_id sai",
     "Backend trả 404 hoặc 422",
     "High"),

    ("TC046", "M07 Guestbook & Collection", "AUTHORIZATION",
     "Guest gửi guestbook",
     "Chưa đăng nhập",
     "POST guestbook",
     "Backend trả 401",
     "High"),

    ("TC047", "M07 Guestbook & Collection", "AUTHORIZATION",
     "User xóa collection của người khác",
     "Có user A và B",
     "User A đoán artifact/collection thuộc B",
     "Không thể xóa dữ liệu của B",
     "Critical"),

    # M08 - Behavior & Analytics: 6
    ("TC048", "M08 Behavior & Analytics", "NORMAL",
     "Ghi event xem panorama",
     "API event hoạt động",
     "POST event_type=view_panorama",
     "Event được lưu",
     "High"),

    ("TC049", "M08 Behavior & Analytics", "NORMAL",
     "Ghi event mở hotspot",
     "Hotspot tồn tại",
     "POST event_type=open_hotspot",
     "Event được lưu với quan hệ hotspot",
     "High"),

    ("TC050", "M08 Behavior & Analytics", "NORMAL",
     "Ghi event xem hiện vật",
     "Artifact tồn tại",
     "POST event_type=view_artifact",
     "Event được lưu",
     "High"),

    ("TC051", "M08 Behavior & Analytics", "VALIDATION",
     "Event thiếu session_id",
     "API hoạt động",
     "POST event không có session_id",
     "Backend trả 422",
     "High"),

    ("TC052", "M08 Behavior & Analytics", "VALIDATION",
     "Event có event_type rỗng",
     "API hoạt động",
     "POST event_type rỗng",
     "Backend trả 422",
     "Medium"),

    ("TC053", "M08 Behavior & Analytics", "AUTHORIZATION",
     "Gọi endpoint thống kê quản trị bằng user thường",
     "User thường đăng nhập",
     "Gọi endpoint analytics quản trị",
     "Trả 403",
     "Critical"),

    # M09 - Admin & Data Integrity: 7
    ("TC054", "M09 Admin & Data Integrity", "NORMAL",
     "Admin xem dữ liệu museum spaces",
     "Admin đăng nhập",
     "Mở trang quản trị spaces",
     "Dữ liệu hiển thị đầy đủ",
     "High"),

    ("TC055", "M09 Admin & Data Integrity", "NORMAL",
     "Admin xem dữ liệu artifacts",
     "Admin đăng nhập",
     "Mở quản trị artifacts",
     "Dữ liệu hiển thị",
     "High"),

    ("TC056", "M09 Admin & Data Integrity", "NORMAL",
     "Admin kiểm tra dữ liệu behavior events",
     "Admin đăng nhập",
     "Mở thống kê hành vi",
     "Dữ liệu được hiển thị theo thời gian",
     "High"),

    ("TC057", "M09 Admin & Data Integrity", "VALIDATION",
     "Tạo artifact thiếu museum_space_id",
     "Admin đăng nhập",
     "Gửi dữ liệu thiếu FK bắt buộc",
     "Backend từ chối",
     "Critical"),

    ("TC058", "M09 Admin & Data Integrity", "VALIDATION",
     "Tạo tour stop thiếu tour_id",
     "Admin đăng nhập",
     "Gửi dữ liệu không đủ trường",
     "Backend trả 422",
     "High"),

    ("TC059", "M09 Admin & Data Integrity", "AUTHORIZATION",
     "User thường truy cập dashboard admin",
     "User thường đăng nhập",
     "Mở trực tiếp đường dẫn quản trị",
     "Trả 403",
     "Critical"),

    ("TC060", "M09 Admin & Data Integrity", "AUTHORIZATION",
     "User thường truy cập endpoint quản trị qua API",
     "User thường đăng nhập",
     "Gọi trực tiếp API admin",
     "Trả 403; không trả dữ liệu quản trị",
     "Critical"),
]


if len(cases) != 60:
    raise RuntimeError(
        f"Test plan phai co 60 ca, hien tai co {len(cases)}."
    )


type_counts = {}

for case in cases:
    case_type = case[2]
    type_counts[case_type] = type_counts.get(case_type, 0) + 1


if type_counts.get("NORMAL", 0) == 0:
    raise RuntimeError("Thieu kich ban NORMAL.")

if type_counts.get("VALIDATION", 0) == 0:
    raise RuntimeError("Thieu kich ban VALIDATION.")

if type_counts.get("AUTHORIZATION", 0) == 0:
    raise RuntimeError("Thieu kich ban AUTHORIZATION.")


workbook = Workbook()

sheet = workbook.active
sheet.title = "CHECKLIST_KIEMTHU"

headers = [
    "Test ID",
    "Module",
    "Scenario Type",
    "Scenario",
    "Precondition",
    "Steps",
    "Expected Result",
    "Priority",
    "Result",
    "Actual Result",
    "Evidence",
    "Browser",
    "Screen Size",
]


sheet.append(headers)

for case in cases:
    sheet.append(
        [
            case[0],
            case[1],
            case[2],
            case[3],
            case[4],
            case[5],
            case[6],
            case[7],
            "",
            "",
            "",
            "Chrome",
            "Desktop 1920x1080",
        ]
    )


header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

header_font = Font(
    bold=True,
    color="FFFFFF",
)

for cell in sheet[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


for row in sheet.iter_rows(
    min_row=2,
    max_row=sheet.max_row,
):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


widths = {
    "A": 12,
    "B": 32,
    "C": 18,
    "D": 42,
    "E": 34,
    "F": 44,
    "G": 48,
    "H": 12,
    "I": 15,
    "J": 35,
    "K": 35,
    "L": 20,
    "M": 22,
}

for column, width in widths.items():
    sheet.column_dimensions[column].width = width


sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions


summary = workbook.create_sheet("TONG_HOP_KIEMTHU")

summary_headers = [
    "Module",
    "So ca",
    "Normal",
    "Validation",
    "Authorization",
]

summary.append(summary_headers)

modules = [
    "M01 Authentication & Roles",
    "M02 Museum Spaces & Panoramas",
    "M03 Artifacts & Hotspots",
    "M04 360 Scene Transitions",
    "M05 Multilingual Audio",
    "M06 Themed Tours & Modes",
    "M07 Guestbook & Collection",
    "M08 Behavior & Analytics",
    "M09 Admin & Data Integrity",
]

for module in modules:
    module_cases = [
        case for case in cases if case[1] == module
    ]

    normal = sum(
        1 for case in module_cases if case[2] == "NORMAL"
    )

    validation = sum(
        1 for case in module_cases if case[2] == "VALIDATION"
    )

    authorization = sum(
        1 for case in module_cases
        if case[2] == "AUTHORIZATION"
    )

    summary.append(
        [
            module,
            len(module_cases),
            normal,
            validation,
            authorization,
        ]
    )


summary.append(
    [
        "TONG",
        f"=SUM(B2:B{summary.max_row})",
        f"=SUM(C2:C{summary.max_row})",
        f"=SUM(D2:D{summary.max_row})",
        f"=SUM(E2:E{summary.max_row})",
    ]
)


for cell in summary[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


for row in summary.iter_rows(
    min_row=2,
    max_row=summary.max_row,
):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


summary.column_dimensions["A"].width = 36
summary.column_dimensions["B"].width = 12
summary.column_dimensions["C"].width = 12
summary.column_dimensions["D"].width = 15
summary.column_dimensions["E"].width = 18

summary.freeze_panes = "A2"


workbook.save(OUTPUT_FILE)


print("==============================================")
print("DT-16 — TEST PLAN GENERATOR")
print("==============================================")
print(f"Output : {OUTPUT_FILE}")
print(f"Total  : {len(cases)}")
print(f"NORMAL : {type_counts.get('NORMAL', 0)}")
print(f"VALID  : {type_counts.get('VALIDATION', 0)}")
print(f"AUTHZ  : {type_counts.get('AUTHORIZATION', 0)}")
print("PASS   : TEST PLAN CO 60 CA")
